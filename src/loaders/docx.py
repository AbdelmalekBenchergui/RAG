import docx
import io


def parse_docx_bytes(file_bytes: bytes, filename: str):
    doc = docx.Document(io.BytesIO(file_bytes))
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            full_text.append(" | ".join(row_text))
    return "\n\n".join(full_text), {"filename": filename, "type": "docx"}
